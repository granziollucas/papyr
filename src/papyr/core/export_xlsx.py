"""XLSX export."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from papyr.core.export_csv import CSV_COLUMNS, _csv_value
from papyr.core.models import PaperRecord


def _sheet_name_for_origin(origin: str) -> str:
    key = (origin or "").strip().lower()
    if key == "crossref":
        return "Crossref"
    if key in {"arxiv", "arxiv.org"}:
        return "arXiv"
    if key == "ssrn":
        return "SSRN"
    return origin.strip() or "Other"


def _ensure_header(sheet) -> None:
    if sheet.max_row == 1 and sheet.max_column == 1 and sheet.cell(row=1, column=1).value is None:
        sheet.append(CSV_COLUMNS)
        return
    first_row = [cell.value for cell in sheet[1]]
    if first_row != CSV_COLUMNS:
        return


def export_xlsx(records: list[PaperRecord], path: Path, append: bool = False) -> None:
    """Write XLSX with one sheet per provider."""
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:  # pragma: no cover - handled by runtime environment
        raise RuntimeError("openpyxl is required for XLSX export") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    default_sheet = None
    if append and path.exists():
        workbook = load_workbook(path)
    else:
        workbook = Workbook()
        default_sheet = workbook.active

    grouped: dict[str, list[PaperRecord]] = defaultdict(list)
    for record in records:
        grouped[_sheet_name_for_origin(record.origin)].append(record)

    for sheet_name, items in grouped.items():
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.create_sheet(sheet_name)
        _ensure_header(sheet)
        for record in items:
            sheet.append(
                [
                    _csv_value(record.authors),
                    _csv_value(record.title),
                    _csv_value(record.abstract),
                    _csv_value(record.origin),
                    _csv_value(record.volume),
                    _csv_value(record.issue),
                    _csv_value(record.pages),
                    _csv_value(record.publisher),
                    _csv_value(record.month),
                    _csv_value(record.year),
                    _csv_value(record.type),
                    _csv_value(record.keywords),
                    _csv_value(record.citations),
                    _csv_value(record.oa),
                    _csv_value(record.id),
                    _csv_value(record.url),
                    _csv_value(record.license),
                    _csv_value(record.retrieved_at),
                    _csv_value(record.query_hash),
                    _csv_value(record.duplicate_of),
                ]
            )

    if grouped and default_sheet is not None:
        if (
            default_sheet.max_row == 1
            and default_sheet.max_column == 1
            and default_sheet.cell(row=1, column=1).value is None
        ):
            workbook.remove(default_sheet)

    workbook.save(path)
